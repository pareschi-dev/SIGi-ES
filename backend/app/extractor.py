"""Deterministic, bounded document parsing that emits unconfirmed candidates only."""

from dataclasses import dataclass
import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata
from zipfile import BadZipFile, ZipFile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session
import xlrd

from app.business_catalog import MATERIALS, ORGANIZATIONS, SERVICES, SUPPLIERS
from app.models import ExtractionCandidate as ExtractionCandidateRecord
from app.models import SourceDocument

MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_WORKSHEETS = 100
MAX_SPREADSHEET_CELLS = 50_000
MAX_OFFICE_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
MAX_ZIP_ENTRIES = 10_000
MAX_COMPRESSION_RATIO = 200
MAX_CANDIDATES = 1_000

SEI_PATTERN = re.compile(r"(?<!\d)(\d{5}\.\d{6}/\d{4}-\d{2})(?!\d)")
BRL_PATTERN = re.compile(r"R\$\s*((?:\d{1,3}(?:\.\d{3})+|\d+),\d{2})", re.IGNORECASE)
DUE_DATE_PATTERN = re.compile(
    r"(?:vencimento|vence\s+em|pagar\s+at[eé]|data\s+de\s+vencimento)"
    r"\D{0,32}(\d{2}/\d{2}/\d{4})",
    re.IGNORECASE,
)
RECEIPT_TERM_PATTERN = re.compile(
    r"\b(?:TERMO\s+DE\s+RECEBIMENTO(?:\s+(?:PROVIS[ÓO]RIO|DEFINITIVO))?|"
    r"TERMO\s+DE\s+ATESTO|ATESTO\s+DE\s+RECEBIMENTO)\b",
    re.IGNORECASE,
)
INVOICE_CONTEXT_PATTERN = re.compile(
    r"\b(?:FATURA|NOTA\s+FISCAL|NF\s*[-º°]?\s*E|DOCUMENTO\s+FISCAL)\b",
    re.IGNORECASE,
)
AMOUNT_FIELD_PRECEDENCE = (
    "reimbursement_term_amount_brl",
    "invoice_amount_brl",
    "amount_brl",
)


@dataclass(frozen=True, slots=True)
class FieldCandidate:
    """A candidate value and its evidence pointer; never a confirmed business fact."""

    field_name: str
    raw_value: str
    normalized_value: str | None
    extraction_method: str
    evidence_location: str
    confidence: Decimal | None = None


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    """Candidates and parser status for one document."""

    parser: str
    candidates: tuple[FieldCandidate, ...]
    warnings: tuple[str, ...]


def _check_file_size(path: Path, max_file_bytes: int) -> None:
    if max_file_bytes < 1:
        raise ValueError("max_file_bytes deve ser positivo")
    if path.stat().st_size > max_file_bytes:
        raise ValueError("file_too_large")


def validate_office_archive(path: Path) -> None:
    """Bound ZIP metadata before parsing XLSX/DOCX without extracting members."""
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise ValueError("office_archive_entry_limit_exceeded")
            uncompressed_total = 0
            for entry in entries:
                member = Path(entry.filename)
                if member.is_absolute() or ".." in member.parts:
                    raise ValueError("unsafe_office_archive_member")
                uncompressed_total += entry.file_size
                if uncompressed_total > MAX_OFFICE_UNCOMPRESSED_BYTES:
                    raise ValueError("office_archive_expansion_limit_exceeded")
                if entry.compress_size and entry.file_size / entry.compress_size > MAX_COMPRESSION_RATIO:
                    raise ValueError("office_archive_compression_ratio_exceeded")
    except BadZipFile as exc:
        raise ValueError("invalid_office_archive") from exc


def _normalize_brl(raw_value: str) -> str | None:
    try:
        normalized = Decimal(raw_value.replace(".", "").replace(",", "."))
    except InvalidOperation:
        return None
    if not normalized.is_finite() or normalized <= 0:
        return None
    return format(normalized.quantize(Decimal("0.01")), "f")


def _normalize_catalog_text(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value).casefold()
    without_marks = "".join(char for char in decomposed if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", without_marks).strip()


def _catalog_matches(
    text: str,
    *,
    field_name: str,
    items: list[dict[str, object]],
    labels: tuple[str, ...],
    evidence_location: str,
    extraction_method: str,
) -> tuple[FieldCandidate, ...]:
    normalized_text = _normalize_catalog_text(text)
    found: list[FieldCandidate] = []
    for item in items:
        matched_label = next(
            (
                label
                for label in labels
                if label in item
                and re.search(
                    rf"(?<![a-z0-9]){re.escape(_normalize_catalog_text(str(item[label])))}(?![a-z0-9])",
                    normalized_text,
                )
            ),
            None,
        )
        if matched_label is None:
            continue
        found.append(
            FieldCandidate(
                field_name=field_name,
                raw_value=str(item[matched_label]),
                normalized_value=str(item["id"]),
                extraction_method=f"user_catalog:{extraction_method[:12]}",
                evidence_location=evidence_location,
            )
        )
    return tuple(found)


def _amount_field_name(
    text: str,
    match_start: int,
    match_end: int,
    document_context: str | None = None,
) -> str:
    """Classify an amount only when nearby/page evidence identifies its source."""
    has_receipt_term = RECEIPT_TERM_PATTERN.search(text) is not None
    has_invoice_context = INVOICE_CONTEXT_PATTERN.search(text) is not None
    has_document_receipt_term = bool(
        document_context and RECEIPT_TERM_PATTERN.search(document_context)
    )
    has_document_invoice_context = bool(
        document_context and INVOICE_CONTEXT_PATTERN.search(document_context)
    )

    line_start = text.rfind("\n", 0, match_start) + 1
    line_end = text.find("\n", match_end)
    if line_end < 0:
        line_end = len(text)
    previous_start = text.rfind("\n", 0, max(0, line_start - 1)) + 1
    next_end = text.find("\n", line_end + 1)
    if next_end < 0:
        next_end = len(text)
    local_context = text[previous_start:next_end]
    local_has_receipt_term = RECEIPT_TERM_PATTERN.search(local_context) is not None
    local_has_invoice_context = INVOICE_CONTEXT_PATTERN.search(local_context) is not None

    if local_has_receipt_term:
        return "reimbursement_term_amount_brl"
    if local_has_invoice_context:
        return "invoice_amount_brl"
    if has_receipt_term and not has_invoice_context:
        return "reimbursement_term_amount_brl"
    if has_invoice_context and not has_receipt_term:
        return "invoice_amount_brl"
    if has_document_receipt_term and not has_document_invoice_context:
        return "reimbursement_term_amount_brl"
    if has_document_invoice_context and not has_document_receipt_term:
        return "invoice_amount_brl"
    return "amount_brl"


def preferred_amount_candidates(candidates):
    """Return best-source amount candidates without resolving multiple values."""
    amount_candidates = [
        candidate for candidate in candidates if candidate.field_name in AMOUNT_FIELD_PRECEDENCE
    ]
    for field_name in AMOUNT_FIELD_PRECEDENCE:
        preferred = [candidate for candidate in amount_candidates if candidate.field_name == field_name]
        if preferred:
            return tuple(preferred)
    return ()


def extract_candidates_from_text(
    text: str,
    *,
    evidence_location: str,
    extraction_method: str,
    document_context: str | None = None,
    max_candidates: int = MAX_CANDIDATES,
) -> tuple[FieldCandidate, ...]:
    """Extract SEI-shaped IDs, labeled BRL amounts, and labeled due-date candidates."""
    if max_candidates < 1:
        raise ValueError("max_candidates deve ser positivo")
    candidates: list[FieldCandidate] = []
    seen: set[tuple[str, str, str]] = set()

    def add(field_name: str, raw_value: str, normalized: str | None) -> None:
        key = (field_name, raw_value, evidence_location)
        if key in seen:
            return
        if len(candidates) >= max_candidates:
            raise ValueError("candidate_limit_exceeded")
        seen.add(key)
        candidates.append(
            FieldCandidate(
                field_name=field_name,
                raw_value=raw_value,
                normalized_value=normalized,
                extraction_method=extraction_method,
                evidence_location=evidence_location,
            )
        )

    for match in SEI_PATTERN.finditer(text):
        add("process_number", match.group(1), match.group(1))
    for match in BRL_PATTERN.finditer(text):
        raw_value = match.group(1)
        normalized = _normalize_brl(raw_value)
        if normalized is not None:
            add(
                _amount_field_name(text, match.start(), match.end(), document_context),
                raw_value,
                normalized,
            )
    for match in DUE_DATE_PATTERN.finditer(text):
        raw_value = match.group(1)
        try:
            normalized = datetime.strptime(raw_value, "%d/%m/%Y").date().isoformat()
        except ValueError:
            continue
        add("due_date", raw_value, normalized)
    for catalog_candidate in (
        *_catalog_matches(
            text,
            field_name="organization_reference",
            items=ORGANIZATIONS,
            labels=("nome", "sigla"),
            evidence_location=evidence_location,
            extraction_method=extraction_method,
        ),
        *_catalog_matches(
            text,
            field_name="material_reference",
            items=MATERIALS,
            labels=("nome",),
            evidence_location=evidence_location,
            extraction_method=extraction_method,
        ),
        *_catalog_matches(
            text,
            field_name="service_reference",
            items=SERVICES,
            labels=("nome",),
            evidence_location=evidence_location,
            extraction_method=extraction_method,
        ),
        *_catalog_matches(
            text,
            field_name="supplier_reference",
            items=SUPPLIERS,
            labels=("nome",),
            evidence_location=evidence_location,
            extraction_method=extraction_method,
        ),
    ):
        add(
            catalog_candidate.field_name,
            catalog_candidate.raw_value,
            catalog_candidate.normalized_value,
        )
    return tuple(candidates)


def extract_document(
    path: Path,
    *,
    max_file_bytes: int = MAX_FILE_BYTES,
    max_pdf_pages: int = MAX_PDF_PAGES,
    max_spreadsheet_cells: int = MAX_SPREADSHEET_CELLS,
) -> ExtractionResult:
    """Read supported document content and return candidates with page/cell pointers.

    PDF parsing is native-text only. XLSX/XLSM formulas are not evaluated or used;
    workbooks are opened read-only and are never saved. Legacy XLS exposes cached
    cell values; formula provenance may not be available, so every result remains
    an unconfirmed candidate. CSV is read as plain text. OCR and external services
    are deliberately not invoked.
    """
    if max_pdf_pages < 1 or max_spreadsheet_cells < 1:
        raise ValueError("page_and_cell_limits_must_be_positive")
    source = Path(path).resolve(strict=True)
    if not source.is_file():
        raise ValueError("source_must_be_a_regular_file")
    _check_file_size(source, max_file_bytes)
    extension = source.suffix.lower()
    candidates: list[FieldCandidate] = []
    warnings: list[str] = []

    if extension == ".pdf":
        reader = PdfReader(source, strict=False)
        if reader.is_encrypted:
            raise ValueError("encrypted_pdf_requires_authorized_manual_handling")
        if len(reader.pages) > max_pdf_pages:
            raise ValueError("pdf_page_limit_exceeded")
        for page_index, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                candidates.extend(
                    extract_candidates_from_text(
                        page_text,
                        evidence_location=f"page:{page_index}",
                        extraction_method="pdf_native_text",
                        max_candidates=MAX_CANDIDATES - len(candidates),
                    )
                )
        if not candidates:
            warnings.append("no_native_text_or_matching_candidates;_ocr_not_attempted")
        return ExtractionResult("pdf_native_text", tuple(candidates), tuple(warnings))

    if extension in {".xlsx", ".xlsm"}:
        validate_office_archive(source)
        workbook = load_workbook(
            source, read_only=True, data_only=False, keep_vba=False, keep_links=False
        )
        try:
            if len(workbook.worksheets) > MAX_WORKSHEETS:
                raise ValueError("worksheet_limit_exceeded")
            processed_cells = 0
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    row_values: list[str] = []
                    for cell in row:
                        processed_cells += 1
                        if processed_cells > max_spreadsheet_cells:
                            raise ValueError("spreadsheet_cell_limit_exceeded")
                        value = cell.value
                        if value is None or cell.data_type == "f":
                            continue
                        raw_value = str(value).strip()
                        if not raw_value:
                            continue
                        row_values.append(raw_value)
                        candidates.extend(
                            extract_candidates_from_text(
                                raw_value,
                                evidence_location=f"sheet:{worksheet.title}!{cell.coordinate}",
                                extraction_method="spreadsheet_cell",
                                max_candidates=MAX_CANDIDATES - len(candidates),
                            )
                        )
                    if row_values:
                        candidates.extend(
                            extract_candidates_from_text(
                                " | ".join(row_values),
                                evidence_location=f"sheet:{worksheet.title}!row:{row[0].row if row else 0}",
                                extraction_method="spreadsheet_row",
                                max_candidates=MAX_CANDIDATES - len(candidates),
                            )
                        )
        finally:
            workbook.close()
        if not candidates:
            warnings.append("no_matching_candidates_found")
        return ExtractionResult("spreadsheet_cells", tuple(candidates), tuple(warnings))

    if extension == ".xls":
        workbook = xlrd.open_workbook(source, on_demand=True, formatting_info=False)
        try:
            if workbook.nsheets > MAX_WORKSHEETS:
                raise ValueError("worksheet_limit_exceeded")
            processed_cells = 0
            for worksheet in workbook.sheets():
                for row_index in range(worksheet.nrows):
                    row_values: list[str] = []
                    for column_index in range(worksheet.ncols):
                        processed_cells += 1
                        if processed_cells > max_spreadsheet_cells:
                            raise ValueError("spreadsheet_cell_limit_exceeded")
                        cell = worksheet.cell(row_index, column_index)
                        if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
                            continue
                        raw_value = str(cell.value).strip()
                        if not raw_value:
                            continue
                        row_values.append(raw_value)
                        candidates.extend(
                            extract_candidates_from_text(
                                raw_value,
                                evidence_location=(
                                    f"sheet:{worksheet.name}!row:{row_index + 1},"
                                    f"column:{column_index + 1}"
                                ),
                                extraction_method="xls_cell_value_unverified_formula_origin",
                                max_candidates=MAX_CANDIDATES - len(candidates),
                            )
                        )
                    if row_values:
                        candidates.extend(
                            extract_candidates_from_text(
                                " | ".join(row_values),
                                evidence_location=f"sheet:{worksheet.name}!row:{row_index + 1}",
                                extraction_method="xls_row_cached_values_unverified_formula_origin",
                                max_candidates=MAX_CANDIDATES - len(candidates),
                            )
                        )
        finally:
            workbook.release_resources()
        warnings.append("legacy_xls_formula_origin_unverified")
        if not candidates:
            warnings.append("no_matching_candidates_found")
        return ExtractionResult("legacy_xls_cells", tuple(candidates), tuple(warnings))

    if extension == ".csv":
        raw_bytes = source.read_bytes()
        try:
            text_content = raw_bytes.decode("utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            text_content = raw_bytes.decode("cp1252")
            encoding = "cp1252"
        sample = text_content[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
        processed_cells = 0
        for row_index, row in enumerate(csv.reader(text_content.splitlines(), dialect), start=1):
            row_values: list[str] = []
            for column_index, value in enumerate(row, start=1):
                processed_cells += 1
                if processed_cells > max_spreadsheet_cells:
                    raise ValueError("spreadsheet_cell_limit_exceeded")
                value = value.strip()
                if not value:
                    continue
                row_values.append(value)
                candidates.extend(
                    extract_candidates_from_text(
                        value,
                        evidence_location=f"row:{row_index},column:{column_index}",
                        extraction_method=f"csv_cell:{encoding}",
                        max_candidates=MAX_CANDIDATES - len(candidates),
                    )
                )
            if row_values:
                candidates.extend(
                    extract_candidates_from_text(
                        " | ".join(row_values),
                        evidence_location=f"row:{row_index}",
                        extraction_method=f"csv_row:{encoding}",
                        max_candidates=MAX_CANDIDATES - len(candidates),
                    )
                )
        if not candidates:
            warnings.append("no_matching_candidates_found")
        return ExtractionResult("csv_cells", tuple(candidates), tuple(warnings))

    if extension == ".docx":
        validate_office_archive(source)
        document = Document(source)
        text_blocks: list[tuple[str, str]] = []
        text_blocks.extend(
            (paragraph.text, f"paragraph:{index}")
            for index, paragraph in enumerate(document.paragraphs, start=1)
        )
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                for column_index, cell in enumerate(row.cells, start=1):
                    text_blocks.append(
                        (cell.text, f"table:{table_index},row:{row_index},column:{column_index}")
                    )
        document_context = "\n".join(text for text, _ in text_blocks)
        for text, location in text_blocks:
            candidates.extend(
                extract_candidates_from_text(
                    text,
                    evidence_location=location,
                    extraction_method="docx_text",
                    document_context=document_context,
                    max_candidates=MAX_CANDIDATES - len(candidates),
                )
            )
        if not candidates:
            warnings.append("no_matching_candidates_found")
        return ExtractionResult("docx_text", tuple(candidates), tuple(warnings))

    raise ValueError("unsupported_document_type")


def persist_candidates(
    session: Session,
    document_id,
    candidates: tuple[FieldCandidate, ...],
) -> int:
    """Persist candidates idempotently without promoting them to confirmed fields."""
    document = session.get(SourceDocument, document_id)
    if document is None:
        raise ValueError("source_document_not_found")
    inserted = 0
    seen_keys: set[tuple[str, str, str, str]] = set()
    for candidate in candidates:
        key = (
            candidate.field_name,
            candidate.raw_value,
            candidate.evidence_location,
            candidate.extraction_method,
        )
        if key in seen_keys:
            continue
        seen_keys.add(key)
        statement = select(ExtractionCandidateRecord).where(
            ExtractionCandidateRecord.document_id == document_id,
            ExtractionCandidateRecord.field_name == candidate.field_name,
            ExtractionCandidateRecord.raw_value == candidate.raw_value,
            ExtractionCandidateRecord.evidence_location == candidate.evidence_location,
            ExtractionCandidateRecord.extraction_method == candidate.extraction_method,
        )
        if session.scalar(statement) is not None:
            continue
        session.add(
            ExtractionCandidateRecord(
                document_id=document_id,
                field_name=candidate.field_name,
                raw_value=candidate.raw_value,
                normalized_value=candidate.normalized_value,
                extraction_method=candidate.extraction_method,
                evidence_location=candidate.evidence_location,
                confidence=None,
                review_state="candidate",
            )
        )
        inserted += 1
    session.commit()
    return inserted
