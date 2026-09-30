"""Read-only workbook preview using an explicit, caller-supplied column mapping."""

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from app.extractor import MAX_FILE_BYTES, validate_office_archive

MAX_PREVIEW_ROWS = 10_000
MAX_AMOUNT_TEXT = re.compile(r"^-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d{1,2})?$")
COLUMN_LETTERS = re.compile(r"^[A-Za-z]{1,3}$")


@dataclass(frozen=True, slots=True)
class WorkbookMapping:
    """Required source worksheet and field-to-column mapping.

    A reference may be an Excel column letter (for example ``C``) or an exact
    header label. Organization, period, and available balance are mandatory.
    """

    sheet_name: str
    header_row: int
    organization: str
    period: str
    available_balance: str
    committed_balance: str | None = None
    paid_amount: str | None = None
    period_format: str = "%m/%Y"
    max_rows: int = MAX_PREVIEW_ROWS


@dataclass(frozen=True, slots=True)
class BalanceRowCandidate:
    """Normalized spreadsheet row candidate with source cell references."""

    row_number: int
    organization_raw: str
    period_start: date
    available_balance: Decimal
    committed_balance: Decimal | None
    paid_amount: Decimal | None
    source_cells: dict[str, str]


@dataclass(frozen=True, slots=True)
class BalanceRowIssue:
    """Row rejection reason without copying potentially sensitive cell contents."""

    row_number: int
    code: str
    field_name: str | None = None


@dataclass(frozen=True, slots=True)
class WorkbookPreview:
    """Preview results; the workbook is never modified or persisted by this module."""

    source_sha256: str
    sheet_name: str
    rows_examined: int
    accepted: tuple[BalanceRowCandidate, ...]
    issues: tuple[BalanceRowIssue, ...]


def _column_index(reference: str) -> int | None:
    if not COLUMN_LETTERS.fullmatch(reference):
        return None
    index = 0
    for character in reference.upper():
        index = index * 26 + ord(character) - ord("A") + 1
    if index > 16_384:
        raise ValueError("column_reference_out_of_range")
    return index - 1


def _resolve_column(reference: str, headers: tuple[Any, ...]) -> int:
    if not reference.strip():
        raise ValueError("empty_column_reference")
    letter_index = _column_index(reference.strip())
    if letter_index is not None:
        return letter_index
    matches = [
        index for index, value in enumerate(headers)
        if value is not None and str(value).strip().casefold() == reference.strip().casefold()
    ]
    if not matches:
        raise ValueError("configured_header_not_found")
    if len(matches) > 1:
        raise ValueError("configured_header_is_ambiguous")
    return matches[0]


def _parse_period(value: Any, period_format: str) -> date | None:
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    if value is None:
        return None
    try:
        parsed = datetime.strptime(str(value).strip(), period_format).date()
    except ValueError:
        return None
    return date(parsed.year, parsed.month, 1)


def _parse_amount(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        amount = value
    elif isinstance(value, (int, float)):
        try:
            amount = Decimal(str(value))
        except InvalidOperation:
            return None
    else:
        text = str(value).replace("R$", "").replace("\u00a0", " ").strip()
        if not MAX_AMOUNT_TEXT.fullmatch(text):
            return None
        text = text.replace(".", "").replace(",", ".")
        try:
            amount = Decimal(text)
        except InvalidOperation:
            return None
    if not amount.is_finite() or amount.as_tuple().exponent < -2:
        return None
    return amount.quantize(Decimal("0.01"))


def _mapped_cell(row: tuple[Cell, ...], index: int) -> Cell | None:
    return row[index] if index < len(row) else None


def preview_workbook(
    path: Path,
    mapping: WorkbookMapping,
    *,
    max_file_bytes: int = MAX_FILE_BYTES,
) -> WorkbookPreview:
    """Read a configured Excel sheet and validate rows without importing them.

    Formulas are ignored, macros are not executed, no default sheet or column
    mapping is assumed, and errors do not include raw financial cell values.
    """
    if not mapping.sheet_name.strip():
        raise ValueError("sheet_name_required")
    if mapping.header_row < 1 or mapping.max_rows < 1 or mapping.max_rows > MAX_PREVIEW_ROWS:
        raise ValueError("invalid_preview_limits")
    if max_file_bytes < 1:
        raise ValueError("max_file_bytes_must_be_positive")

    source = Path(path).resolve(strict=True)
    if not source.is_file() or source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("source_must_be_xlsx_or_xlsm_file")
    initial_stat = source.stat()
    if initial_stat.st_size > max_file_bytes:
        raise ValueError("file_too_large")
    validate_office_archive(source)

    workbook = load_workbook(
        source, read_only=True, data_only=False, keep_vba=False, keep_links=False
    )
    accepted: list[BalanceRowCandidate] = []
    issues: list[BalanceRowIssue] = []
    rows_examined = 0
    try:
        if mapping.sheet_name not in workbook.sheetnames:
            raise ValueError("configured_sheet_not_found")
        worksheet = workbook[mapping.sheet_name]
        try:
            header_cells = next(
                worksheet.iter_rows(min_row=mapping.header_row, max_row=mapping.header_row)
            )
        except StopIteration as exc:
            raise ValueError("configured_header_row_not_found") from exc
        headers = tuple(cell.value for cell in header_cells)

        references = {
            "organization": mapping.organization,
            "period": mapping.period,
            "available_balance": mapping.available_balance,
        }
        if mapping.committed_balance:
            references["committed_balance"] = mapping.committed_balance
        if mapping.paid_amount:
            references["paid_amount"] = mapping.paid_amount
        indices = {field: _resolve_column(reference, headers) for field, reference in references.items()}
        data_rows = worksheet.iter_rows(min_row=mapping.header_row + 1)
        for row in data_rows:
            row_number = row[0].row if row else mapping.header_row + rows_examined + 1
            if not any(cell.value is not None for cell in row):
                continue
            rows_examined += 1
            if rows_examined > mapping.max_rows:
                raise ValueError("preview_row_limit_exceeded")

            selected = {field: _mapped_cell(row, index) for field, index in indices.items()}
            formula_field = next(
                (field for field, cell in selected.items() if cell is not None and cell.data_type == "f"),
                None,
            )
            if formula_field:
                issues.append(BalanceRowIssue(row_number, "formula_not_imported", formula_field))
                continue

            organization_value = selected["organization"].value if selected["organization"] else None
            period_value = selected["period"].value if selected["period"] else None
            available_value = selected["available_balance"].value if selected["available_balance"] else None
            organization = str(organization_value).strip() if organization_value is not None else ""
            period_start = _parse_period(period_value, mapping.period_format)
            available = _parse_amount(available_value)

            if not organization:
                issues.append(BalanceRowIssue(row_number, "missing_organization", "organization"))
                continue
            if period_start is None:
                issues.append(BalanceRowIssue(row_number, "invalid_period", "period"))
                continue
            if available is None:
                issues.append(BalanceRowIssue(row_number, "invalid_available_balance", "available_balance"))
                continue

            optional_amounts: dict[str, Decimal | None] = {}
            invalid_optional = False
            for field in ("committed_balance", "paid_amount"):
                cell = selected.get(field)
                if cell is None or cell.value is None:
                    optional_amounts[field] = None
                    continue
                amount = _parse_amount(cell.value)
                if amount is None:
                    issues.append(BalanceRowIssue(row_number, "invalid_optional_amount", field))
                    invalid_optional = True
                    break
                optional_amounts[field] = amount
            if invalid_optional:
                continue

            source_cells = {
                field: cell.coordinate
                for field, cell in selected.items()
                if cell is not None
            }
            accepted.append(
                BalanceRowCandidate(
                    row_number=row_number,
                    organization_raw=organization,
                    period_start=period_start,
                    available_balance=available,
                    committed_balance=optional_amounts.get("committed_balance"),
                    paid_amount=optional_amounts.get("paid_amount"),
                    source_cells=source_cells,
                )
            )
    finally:
        workbook.close()

    final_stat = source.stat()
    if (
        initial_stat.st_size != final_stat.st_size
        or initial_stat.st_mtime_ns != final_stat.st_mtime_ns
        or initial_stat.st_ctime_ns != final_stat.st_ctime_ns
    ):
        raise ValueError("file_changed_during_preview")
    digest = hashlib.sha256()
    with source.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return WorkbookPreview(
        source_sha256=digest.hexdigest(),
        sheet_name=mapping.sheet_name,
        rows_examined=rows_examined,
        accepted=tuple(accepted),
        issues=tuple(issues),
    )
