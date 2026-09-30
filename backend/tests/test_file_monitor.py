"""Tests for the opt-in read-only file event monitor."""

from pathlib import Path
from time import monotonic, sleep

from docx import Document
from sqlalchemy import select

from app.file_monitor import DirectoryMonitor
from app.models import ExtractionCandidate, SourceDocument


def test_monitor_indexes_new_supported_file_without_modifying_it(tmp_path: Path, client_and_session) -> None:
    _, session_factory = client_and_session
    root = tmp_path / "monitor-root"
    root.mkdir()
    document = root / "new-receipt.docx"
    source = Document()
    source.add_paragraph("TERMO DE RECEBIMENTO DEFINITIVO")
    source.add_paragraph("Valor recebido: R$ 198,50")
    source.save(document)
    original = document.read_bytes()
    monitor = DirectoryMonitor(root, source_key="monitor-test", session_factory=session_factory)
    monitor.start()
    try:
        from watchdog.events import FileCreatedEvent

        monitor.enqueue_event(FileCreatedEvent(str(document)))
        deadline = monotonic() + 5
        while monotonic() < deadline and monitor.status().parsed == 0:
            sleep(0.05)
        assert monitor.status().indexed == 1
        assert monitor.status().parsed == 1
        assert document.read_bytes() == original
        session = session_factory()
        record = session.scalar(select(SourceDocument))
        amount_candidate = session.scalar(
            select(ExtractionCandidate).where(
                ExtractionCandidate.field_name == "reimbursement_term_amount_brl"
            )
        )
        session.close()
        assert record is not None
        assert record.source_key == "monitor-test"
        assert record.relative_path == "new-receipt.docx"
        assert len(record.sha256) == 64
        assert record.processing_state == "review"
        assert amount_candidate is not None
        assert amount_candidate.normalized_value == "198.50"
    finally:
        monitor.stop()


def test_monitor_ignores_unapproved_extension_and_symlink(tmp_path: Path, client_and_session) -> None:
    _, session_factory = client_and_session
    root = tmp_path / "monitor-root"
    root.mkdir()
    outside = tmp_path / "outside.pdf"
    outside.write_bytes(b"outside")
    unsupported = root / "not-a-document.exe"
    unsupported.write_bytes(b"synthetic executable")
    monitor = DirectoryMonitor(root, source_key="monitor-test", session_factory=session_factory)
    monitor.start()
    try:
        from watchdog.events import FileCreatedEvent

        monitor.enqueue_event(FileCreatedEvent(str(unsupported)))
        link = root / "linked.pdf"
        try:
            link.symlink_to(outside)
        except OSError:
            link = None
        if link:
            monitor.enqueue_event(FileCreatedEvent(str(link)))
        sleep(0.2)
        assert monitor.status().queued == 0
        assert monitor.status().indexed == 0
    finally:
        monitor.stop()
