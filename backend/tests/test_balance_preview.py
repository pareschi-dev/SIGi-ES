"""Preview-only tests using synthetic Excel workbooks."""

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.balance_preview import WorkbookMapping, preview_workbook


def test_preview_maps_headers_normalizes_values_and_keeps_cell_provenance(tmp_path: Path) -> None:
    path = tmp_path / "synthetic-balances.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Saldos sintéticos"
    sheet.append(["Unidade", "Competência", "Saldo disponível", "Empenhado", "Pago"])
    sheet.append(["Órgão Exemplo", "09/2026", "R$ 1.234,56", 200.15, 100])
    sheet.append(["Outra unidade", date(2026, 9, 30), "sem valor", None, None])
    sheet.append(["Fórmula", "09/2026", "=SUM(D2:E2)", 0, 0])
    workbook.save(path)
    workbook.close()
    before = path.read_bytes()

    report = preview_workbook(
        path,
        WorkbookMapping(
            sheet_name="Saldos sintéticos",
            header_row=1,
            organization="Unidade",
            period="Competência",
            available_balance="Saldo disponível",
            committed_balance="Empenhado",
            paid_amount="Pago",
        ),
    )

    assert len(report.accepted) == 1
    row = report.accepted[0]
    assert row.organization_raw == "Órgão Exemplo"
    assert row.period_start == date(2026, 9, 1)
    assert row.available_balance == Decimal("1234.56")
    assert row.committed_balance == Decimal("200.15")
    assert row.paid_amount == Decimal("100.00")
    assert row.source_cells["available_balance"] == "C2"
    assert {issue.code for issue in report.issues} == {"invalid_available_balance", "formula_not_imported"}
    assert report.rows_examined == 3
    assert len(report.source_sha256) == 64
    assert path.read_bytes() == before


def test_preview_supports_explicit_column_letters(tmp_path: Path) -> None:
    path = tmp_path / "letter-map.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Org", "Mês", "Saldo"])
    sheet.append(["Unidade Sintética", "2026-09", 50])
    workbook.save(path)
    workbook.close()

    report = preview_workbook(
        path,
        WorkbookMapping(
            sheet_name="Sheet",
            header_row=1,
            organization="A",
            period="B",
            available_balance="C",
            period_format="%Y-%m",
        ),
    )

    assert report.accepted[0].period_start == date(2026, 9, 1)
    assert report.accepted[0].available_balance == Decimal("50.00")


def test_preview_requires_configured_sheet_and_headers(tmp_path: Path) -> None:
    path = tmp_path / "mapping.xlsx"
    workbook = Workbook()
    workbook.active.append(["Org", "Período", "Saldo"])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="configured_sheet_not_found"):
        preview_workbook(
            path,
            WorkbookMapping("NAO EXISTE", 1, "A", "B", "C"),
        )
    with pytest.raises(ValueError, match="configured_header_not_found"):
        preview_workbook(
            path,
            WorkbookMapping("Sheet", 1, "Órgão", "B", "C"),
        )


def test_preview_rejects_unsupported_file_and_invalid_mapping(tmp_path: Path) -> None:
    path = tmp_path / "not-excel.csv"
    path.write_text("synthetic", encoding="utf-8")

    with pytest.raises(ValueError, match="source_must_be_xlsx_or_xlsm_file"):
        preview_workbook(path, WorkbookMapping("Sheet", 1, "A", "B", "C"))
    with pytest.raises(ValueError, match="invalid_preview_limits"):
        preview_workbook(
            tmp_path / "missing.xlsx",
            WorkbookMapping("Sheet", 1, "A", "B", "C", max_rows=0),
        )
