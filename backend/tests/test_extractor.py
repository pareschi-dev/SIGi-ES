"""Synthetic-document tests for deterministic candidate extraction."""

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
import xlwt
from sqlalchemy import select

from app.extractor import (
    extract_candidates_from_text,
    extract_document,
    persist_candidates,
)
from app.models import ExtractionCandidate, SourceDocument
from app.scanner import index_scan_report, scan_directory


def test_text_extraction_emits_candidates_with_evidence_and_no_confidence() -> None:
    candidates = extract_candidates_from_text(
        "Processo SEI 12345.123456/2026-00; valor total R$ 1.234,56; "
        "vencimento: 31/10/2026.",
        evidence_location="page:2",
        extraction_method="pdf_native_text",
    )

    assert [(item.field_name, item.normalized_value) for item in candidates] == [
        ("process_number", "12345.123456/2026-00"),
        ("amount_brl", "1234.56"),
        ("due_date", "2026-10-31"),
    ]
    assert all(item.evidence_location == "page:2" for item in candidates)
    assert all(item.confidence is None for item in candidates)


def test_invalid_date_and_zero_amount_are_not_candidates() -> None:
    candidates = extract_candidates_from_text(
        "Total R$ 0,00. Vencimento: 31/02/2026.",
        evidence_location="page:1",
        extraction_method="pdf_native_text",
    )

    assert candidates == ()


def test_receipt_term_amount_is_classified_and_preferred_over_invoice_amount() -> None:
    from app.extractor import preferred_amount_candidates

    candidates = extract_candidates_from_text(
        "NOTA FISCAL ELETRÔNICA\nValor total: R$ 210,00\n\n"
        "TERMO DE RECEBIMENTO DEFINITIVO\nValor recebido: R$ 198,50",
        evidence_location="page:1",
        extraction_method="pdf_native_text",
    )

    assert [(candidate.field_name, candidate.normalized_value) for candidate in candidates] == [
        ("invoice_amount_brl", "210.00"),
        ("reimbursement_term_amount_brl", "198.50"),
    ]
    assert [candidate.normalized_value for candidate in preferred_amount_candidates(candidates)] == [
        "198.50"
    ]
    assert all(candidate.confidence is None for candidate in candidates)


def test_ambiguous_amounts_remain_generic_candidates() -> None:
    candidates = extract_candidates_from_text(
        "Valor total: R$ 210,00",
        evidence_location="page:1",
        extraction_method="pdf_native_text",
    )

    assert len(candidates) == 1
    assert candidates[0].field_name == "amount_brl"


def test_text_extraction_matches_user_catalog_entities_with_reference_ids() -> None:
    candidates = extract_candidates_from_text(
        "Unidade SRA/ES; compra de combustível; serviço ENERGIA ELÉTRICA - EDP, fornecedor EDP.",
        evidence_location="page:3",
        extraction_method="pdf_native_text",
    )

    matched = {(candidate.field_name, candidate.normalized_value) for candidate in candidates}
    assert ("organization_reference", "1") in matched
    assert ("material_reference", "4") in matched
    assert ("service_reference", "15") in matched
    assert ("supplier_reference", "9") in matched
    assert all(candidate.evidence_location == "page:3" for candidate in candidates)
    assert all(len(candidate.extraction_method) <= 32 for candidate in candidates)


def test_xlsx_parser_uses_cell_locations_and_skips_formulas(tmp_path: Path) -> None:
    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dados teste"
    sheet["A1"] = "Processo 00000.000001/2026-00"
    sheet["B1"] = "Valor total R$ 3.250,45"
    sheet["C1"] = "Vencimento 30/11/2026"
    sheet["D1"] = '="R$ 999,99"'
    workbook.save(path)
    workbook.close()

    result = extract_document(path)

    assert result.parser == "spreadsheet_cells"
    assert len(result.candidates) == 6
    assert sum(item.extraction_method == "spreadsheet_cell" for item in result.candidates) == 3
    assert sum(item.extraction_method == "spreadsheet_row" for item in result.candidates) == 3
    assert {item.evidence_location for item in result.candidates} == {
        "sheet:Dados teste!A1", "sheet:Dados teste!B1", "sheet:Dados teste!C1",
        "sheet:Dados teste!row:1",
    }
    assert not any(item.normalized_value == "999.99" for item in result.candidates)


def test_docx_parser_tracks_paragraph_evidence(tmp_path: Path) -> None:
    path = tmp_path / "synthetic.docx"
    document = Document()
    document.add_paragraph("Processo 12345.123456/2026-00")
    document.add_paragraph("Pagar até 05/12/2026")
    document.save(path)

    result = extract_document(path)

    assert result.parser == "docx_text"
    assert [item.field_name for item in result.candidates] == ["process_number", "due_date"]
    assert result.candidates[0].evidence_location == "paragraph:1"
    assert result.candidates[1].normalized_value == "2026-12-05"


def test_blank_pdf_reports_ocr_not_attempted(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as stream:
        writer.write(stream)

    result = extract_document(path)

    assert result.parser == "pdf_native_text"
    assert result.candidates == ()
    assert result.warnings == ("no_native_text_or_matching_candidates;_ocr_not_attempted",)


def test_legacy_xls_parser_marks_formula_origin_as_unverified(tmp_path: Path) -> None:
    path = tmp_path / "legacy.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Saldo teste")
    sheet.write(0, 0, "Processo")
    sheet.write(0, 1, "12345.123456/2026-00")
    sheet.write(1, 0, "Valor")
    sheet.write(1, 1, "R$ 125,00")
    workbook.save(str(path))

    result = extract_document(path)

    assert result.parser == "legacy_xls_cells"
    assert any(candidate.field_name == "process_number" for candidate in result.candidates)
    assert "legacy_xls_formula_origin_unverified" in result.warnings
    assert all(candidate.confidence is None for candidate in result.candidates)


def test_csv_parser_extracts_adjacent_label_and_value(tmp_path: Path) -> None:
    path = tmp_path / "legacy.csv"
    path.write_text(
        "Campo;Valor\nProcesso;12345.123456/2026-00\n"
        "Valor;R$ 1.234,56\nVencimento;31/10/2026\n",
        encoding="utf-8",
    )

    result = extract_document(path)

    assert result.parser == "csv_cells"
    assert {candidate.field_name for candidate in result.candidates} >= {
        "process_number", "amount_brl", "due_date"
    }


def test_extraction_candidates_persist_idempotently(client_and_session, tmp_path: Path) -> None:
    _, session_factory = client_and_session
    root = tmp_path / "source"
    root.mkdir()
    file_path = root / "fatura.pdf"
    file_path.write_bytes(b"synthetic source file")
    report = scan_directory(root, "extract-test")
    session = session_factory()
    index_scan_report(session, report)
    source = session.scalar(select(SourceDocument))
    assert source is not None

    candidates = extract_candidates_from_text(
        "Valor R$ 45,67",
        evidence_location="page:1",
        extraction_method="pdf_native_text",
    )
    first_count = persist_candidates(session, source.id, candidates + candidates)
    second_count = persist_candidates(session, source.id, candidates)
    stored = session.scalars(select(ExtractionCandidate)).all()
    session.close()

    assert first_count == 1
    assert second_count == 0
    assert len(stored) == 1
    assert stored[0].review_state == "candidate"
    assert stored[0].confidence is None
    assert stored[0].normalized_value == "45.67"
